#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para gerar banco de questões SAEB baseado nos descritores
Gera arquivo Excel pronto para importação no sistema
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# Banco de questões - Português 5º Ano
QUESTOES_PORT_5 = [
    # D1 - Localizar informações explícitas em um texto
    {
        "codigo": "D1",
        "enunciado": "Leia o texto abaixo:\n\n'Maria acordou cedo naquela manhã de domingo. Ela precisava ir à feira comprar frutas e verduras para a semana. Colocou seu vestido azul favorito e saiu de casa às 7 horas.'\n\nDe acordo com o texto, que dia da semana Maria foi à feira?",
        "alternativa_a": "Segunda-feira",
        "alternativa_b": "Sábado",
        "alternativa_c": "Domingo",
        "alternativa_d": "Sexta-feira",
        "alternativa_e": "Quinta-feira",
        "gabarito": "C"
    },
    {
        "codigo": "D1",
        "enunciado": "Leia o texto:\n\n'O cachorro Rex é muito esperto. Ele sabe sentar, dar a pata e até buscar a bolinha quando seu dono João joga no quintal. Rex tem pelo marrom e olhos castanhos.'\n\nQual é a cor do pelo de Rex?",
        "alternativa_a": "Preto",
        "alternativa_b": "Branco",
        "alternativa_c": "Cinza",
        "alternativa_d": "Marrom",
        "alternativa_e": "Dourado",
        "gabarito": "D"
    },
    # D2 - Estabelecer relações entre partes de um texto
    {
        "codigo": "D2",
        "enunciado": "Leia o texto:\n\n'Pedro estava muito cansado. ELE havia trabalhado o dia inteiro no jardim.'\n\nNo texto, a palavra ELE se refere a:",
        "alternativa_a": "O jardim",
        "alternativa_b": "O dia",
        "alternativa_c": "Pedro",
        "alternativa_d": "O trabalho",
        "alternativa_e": "O cansaço",
        "gabarito": "C"
    },
    {
        "codigo": "D2",
        "enunciado": "'A professora entrou na sala. ELA cumprimentou todos os alunos com um sorriso.'\n\nA palavra ELA substitui:",
        "alternativa_a": "A sala",
        "alternativa_b": "A professora",
        "alternativa_c": "Os alunos",
        "alternativa_d": "O sorriso",
        "alternativa_e": "A escola",
        "gabarito": "B"
    },
    # D3 - Inferir o sentido de uma palavra ou expressão
    {
        "codigo": "D3",
        "enunciado": "Leia a frase:\n\n'O menino ficou radiante quando ganhou o presente de aniversário.'\n\nA palavra RADIANTE significa que o menino ficou:",
        "alternativa_a": "Triste",
        "alternativa_b": "Muito feliz",
        "alternativa_c": "Com raiva",
        "alternativa_d": "Assustado",
        "alternativa_e": "Preocupado",
        "gabarito": "B"
    },
    {
        "codigo": "D3",
        "enunciado": "'Depois da chuva, o arco-íris surgiu majestoso no céu.'\n\nA palavra MAJESTOSO indica que o arco-íris era:",
        "alternativa_a": "Pequeno",
        "alternativa_b": "Escondido",
        "alternativa_c": "Grandioso e bonito",
        "alternativa_d": "Rápido",
        "alternativa_e": "Comum",
        "gabarito": "C"
    },
    # D4 - Inferir uma informação implícita em um texto
    {
        "codigo": "D4",
        "enunciado": "Leia o texto:\n\n'Joãozinho olhou para o céu escuro e apressou o passo. Não queria se molhar.'\n\nPodemos concluir que:",
        "alternativa_a": "Joãozinho estava atrasado para a escola",
        "alternativa_b": "Ia chover",
        "alternativa_c": "Era noite",
        "alternativa_d": "Joãozinho estava correndo",
        "alternativa_e": "O céu estava bonito",
        "gabarito": "B"
    },
    {
        "codigo": "D4",
        "enunciado": "'Ana guardou o guarda-chuva na bolsa e colocou os óculos de sol.'\n\nPodemos concluir que:",
        "alternativa_a": "Estava chovendo",
        "alternativa_b": "Era noite",
        "alternativa_c": "O dia estava ensolarado",
        "alternativa_d": "Ana estava em casa",
        "alternativa_e": "Estava frio",
        "gabarito": "C"
    },
    # D5 - Interpretar texto com auxílio de material gráfico
    {
        "codigo": "D5",
        "enunciado": "Observe a placa:\n\n[SILÊNCIO - HOSPITAL]\n\nEssa placa pede para as pessoas:",
        "alternativa_a": "Correrem",
        "alternativa_b": "Falarem alto",
        "alternativa_c": "Fazerem silêncio",
        "alternativa_d": "Entrarem no hospital",
        "alternativa_e": "Estacionarem",
        "gabarito": "C"
    },
    {
        "codigo": "D5",
        "enunciado": "Observe a placa de trânsito:\n\n[PARE]\n\nEssa placa indica que o motorista deve:",
        "alternativa_a": "Acelerar",
        "alternativa_b": "Parar o veículo",
        "alternativa_c": "Buzinar",
        "alternativa_d": "Virar à direita",
        "alternativa_e": "Estacionar",
        "gabarito": "B"
    },
    # D6 - Identificar o tema de um texto
    {
        "codigo": "D6",
        "enunciado": "Leia o texto:\n\n'A água é essencial para a vida. Devemos economizar água fechando a torneira enquanto escovamos os dentes e tomando banhos mais curtos.'\n\nO tema principal do texto é:",
        "alternativa_a": "Como escovar os dentes",
        "alternativa_b": "A importância de tomar banho",
        "alternativa_c": "A economia de água",
        "alternativa_d": "A torneira do banheiro",
        "alternativa_e": "A vida no planeta",
        "gabarito": "C"
    },
    {
        "codigo": "D6",
        "enunciado": "'Os animais silvestres precisam viver em seu habitat natural. Quando são retirados da natureza, sofrem muito e podem até morrer.'\n\nO texto fala principalmente sobre:",
        "alternativa_a": "Como cuidar de animais domésticos",
        "alternativa_b": "A importância de proteger os animais silvestres",
        "alternativa_c": "Os tipos de animais",
        "alternativa_d": "A morte dos animais",
        "alternativa_e": "Os zoológicos",
        "gabarito": "B"
    },
    # D7 - Identificar o conflito gerador do enredo
    {
        "codigo": "D7",
        "enunciado": "Leia o texto:\n\n'Pedrinho queria muito ir ao parque, mas sua mãe disse que ele só poderia ir depois de fazer a lição de casa. Então, Pedrinho sentou e fez toda a lição rapidamente.'\n\nO problema de Pedrinho era:",
        "alternativa_a": "Não saber fazer a lição",
        "alternativa_b": "Não gostar do parque",
        "alternativa_c": "Precisar fazer a lição antes de ir ao parque",
        "alternativa_d": "Não ter parque perto de casa",
        "alternativa_e": "A mãe não deixá-lo sair",
        "gabarito": "C"
    },
    # D8 - Estabelecer relação causa/consequência
    {
        "codigo": "D8",
        "enunciado": "'Como choveu muito, o rio transbordou.'\n\nA causa do rio ter transbordado foi:",
        "alternativa_a": "O rio ser grande",
        "alternativa_b": "A chuva forte",
        "alternativa_c": "O calor intenso",
        "alternativa_d": "A seca prolongada",
        "alternativa_e": "O vento forte",
        "gabarito": "B"
    },
    {
        "codigo": "D8",
        "enunciado": "'Marina estudou muito para a prova, por isso tirou nota dez.'\n\nMarina tirou nota dez porque:",
        "alternativa_a": "A prova era fácil",
        "alternativa_b": "Ela estudou muito",
        "alternativa_c": "A professora gostava dela",
        "alternativa_d": "Ela copiou do colega",
        "alternativa_e": "Ela teve sorte",
        "gabarito": "B"
    },
    # D9 - Identificar a finalidade de textos
    {
        "codigo": "D9",
        "enunciado": "Leia o texto:\n\n'RECEITA DE BOLO DE CHOCOLATE\nIngredientes: 3 ovos, 2 xícaras de farinha...'\n\nEsse texto serve para:",
        "alternativa_a": "Contar uma história",
        "alternativa_b": "Ensinar a fazer um bolo",
        "alternativa_c": "Vender um produto",
        "alternativa_d": "Dar uma notícia",
        "alternativa_e": "Fazer um convite",
        "gabarito": "B"
    },
    {
        "codigo": "D9",
        "enunciado": "'VENDE-SE: Bicicleta usada, aro 26, cor azul. Preço: R$ 200,00. Telefone: 9999-0000'\n\nA finalidade desse texto é:",
        "alternativa_a": "Contar uma história sobre bicicletas",
        "alternativa_b": "Ensinar a andar de bicicleta",
        "alternativa_c": "Anunciar a venda de uma bicicleta",
        "alternativa_d": "Dar notícias sobre ciclismo",
        "alternativa_e": "Fazer propaganda de uma loja",
        "gabarito": "C"
    },
    # D10 - Identificar marcas linguísticas do locutor e interlocutor
    {
        "codigo": "D10",
        "enunciado": "'Querida vovó, estou com muitas saudades. Quando a senhora vem nos visitar?'\n\nQuem escreveu essa carta?",
        "alternativa_a": "A vovó",
        "alternativa_b": "Um neto ou neta",
        "alternativa_c": "Um vizinho",
        "alternativa_d": "Uma professora",
        "alternativa_e": "Um desconhecido",
        "gabarito": "B"
    },
    # D11 - Distinguir fato de opinião
    {
        "codigo": "D11",
        "enunciado": "Leia as frases:\n\n1. O Brasil tem 26 estados e um Distrito Federal.\n2. O Brasil é o país mais bonito do mundo.\n\nQual frase expressa uma OPINIÃO?",
        "alternativa_a": "Apenas a frase 1",
        "alternativa_b": "Apenas a frase 2",
        "alternativa_c": "As duas frases",
        "alternativa_d": "Nenhuma das frases",
        "alternativa_e": "Não é possível saber",
        "gabarito": "B"
    },
    {
        "codigo": "D11",
        "enunciado": "'A Torre Eiffel fica em Paris e tem 330 metros de altura. É o monumento mais lindo da Europa.'\n\nA OPINIÃO do autor está em:",
        "alternativa_a": "A Torre Eiffel fica em Paris",
        "alternativa_b": "Tem 330 metros de altura",
        "alternativa_c": "É o monumento mais lindo da Europa",
        "alternativa_d": "Todas as informações são fatos",
        "alternativa_e": "Todas as informações são opiniões",
        "gabarito": "C"
    },
    # D12 - Relações lógico-discursivas
    {
        "codigo": "D12",
        "enunciado": "'Ele não foi à escola PORQUE estava doente.'\n\nA palavra PORQUE indica:",
        "alternativa_a": "Tempo",
        "alternativa_b": "Lugar",
        "alternativa_c": "Causa",
        "alternativa_d": "Consequência",
        "alternativa_e": "Dúvida",
        "gabarito": "C"
    },
    # D13 - Identificar efeitos de ironia ou humor
    {
        "codigo": "D13",
        "enunciado": "Leia a tirinha:\n\n- Mãe, tirei zero na prova!\n- Que maravilha, filho!\n\nA mãe disse 'Que maravilha' para:",
        "alternativa_a": "Elogiar o filho",
        "alternativa_b": "Demonstrar ironia/sarcasmo",
        "alternativa_c": "Ficar feliz com a nota",
        "alternativa_d": "Parabenizar o filho",
        "alternativa_e": "Concordar com a nota",
        "gabarito": "B"
    },
    # D14 - Efeito de sentido da pontuação
    {
        "codigo": "D14",
        "enunciado": "Compare as frases:\n\n1. Você vai sair?\n2. Você vai sair!\n\nA diferença entre elas é que:",
        "alternativa_a": "A primeira é uma pergunta e a segunda uma ordem",
        "alternativa_b": "São iguais",
        "alternativa_c": "A segunda é uma pergunta",
        "alternativa_d": "A primeira é uma ordem",
        "alternativa_e": "Ambas são perguntas",
        "gabarito": "A"
    },
    # D15 - Reconhecer diferentes formas de tratar uma informação
    {
        "codigo": "D15",
        "enunciado": "Dois jornais noticiaram o mesmo fato:\n\nJornal A: 'Time local vence por 2 a 1'\nJornal B: 'Time visitante perde partida'\n\nOs dois jornais falam sobre:",
        "alternativa_a": "Jogos diferentes",
        "alternativa_b": "O mesmo jogo, com pontos de vista diferentes",
        "alternativa_c": "Esportes diferentes",
        "alternativa_d": "Notícias falsas",
        "alternativa_e": "Times diferentes",
        "gabarito": "B"
    },
]

# Banco de questões - Matemática 5º Ano
QUESTOES_MAT_5 = [
    # D1 - Localização/movimentação em mapas
    {
        "codigo": "D1",
        "enunciado": "Observe o mapa:\n\n[Mapa com ruas: A escola fica na Rua A, a padaria fica 2 quadras à direita da escola]\n\nPara ir da escola até a padaria, você deve andar:",
        "alternativa_a": "2 quadras à esquerda",
        "alternativa_b": "2 quadras à direita",
        "alternativa_c": "3 quadras em frente",
        "alternativa_d": "1 quadra para trás",
        "alternativa_e": "4 quadras à direita",
        "gabarito": "B"
    },
    {
        "codigo": "D1",
        "enunciado": "No mapa, a casa de João fica na esquina da Rua das Flores com a Avenida Brasil. Para ir até a praça, ele deve virar à direita na Avenida Brasil e andar 3 quadras. A praça fica:",
        "alternativa_a": "À esquerda da casa de João",
        "alternativa_b": "Atrás da casa de João",
        "alternativa_c": "3 quadras à direita na Avenida Brasil",
        "alternativa_d": "Na mesma rua da casa",
        "alternativa_e": "Em outra cidade",
        "gabarito": "C"
    },
    # D2 - Propriedades de poliedros e corpos redondos
    {
        "codigo": "D2",
        "enunciado": "Qual objeto abaixo tem a forma de um CILINDRO?",
        "alternativa_a": "Uma bola de futebol",
        "alternativa_b": "Uma caixa de sapato",
        "alternativa_c": "Uma lata de refrigerante",
        "alternativa_d": "Um cone de sorvete",
        "alternativa_e": "Um dado",
        "gabarito": "C"
    },
    {
        "codigo": "D2",
        "enunciado": "Qual objeto abaixo tem a forma de uma ESFERA?",
        "alternativa_a": "Uma caixa",
        "alternativa_b": "Uma bola de basquete",
        "alternativa_c": "Um rolo de papel",
        "alternativa_d": "Uma pirâmide",
        "alternativa_e": "Um cubo",
        "gabarito": "B"
    },
    # D3 - Figuras bidimensionais
    {
        "codigo": "D3",
        "enunciado": "Quantos lados tem um HEXÁGONO?",
        "alternativa_a": "3 lados",
        "alternativa_b": "4 lados",
        "alternativa_c": "5 lados",
        "alternativa_d": "6 lados",
        "alternativa_e": "8 lados",
        "gabarito": "D"
    },
    {
        "codigo": "D3",
        "enunciado": "Qual figura geométrica plana tem 4 lados iguais e 4 ângulos retos?",
        "alternativa_a": "Triângulo",
        "alternativa_b": "Quadrado",
        "alternativa_c": "Retângulo",
        "alternativa_d": "Losango",
        "alternativa_e": "Trapézio",
        "gabarito": "B"
    },
    # D4 - Quadriláteros
    {
        "codigo": "D4",
        "enunciado": "Qual quadrilátero tem todos os lados paralelos dois a dois?",
        "alternativa_a": "Trapézio",
        "alternativa_b": "Paralelogramo",
        "alternativa_c": "Triângulo",
        "alternativa_d": "Pentágono",
        "alternativa_e": "Círculo",
        "gabarito": "B"
    },
    # D5 - Conservação de medidas em ampliação/redução
    {
        "codigo": "D5",
        "enunciado": "Um quadrado tem lado de 2 cm. Se dobrarmos o tamanho de cada lado, o novo quadrado terá lado de:",
        "alternativa_a": "2 cm",
        "alternativa_b": "3 cm",
        "alternativa_c": "4 cm",
        "alternativa_d": "6 cm",
        "alternativa_e": "8 cm",
        "gabarito": "C"
    },
    # D6 - Estimar medidas
    {
        "codigo": "D6",
        "enunciado": "Qual unidade de medida é mais adequada para medir o comprimento de uma caneta?",
        "alternativa_a": "Quilômetros",
        "alternativa_b": "Metros",
        "alternativa_c": "Centímetros",
        "alternativa_d": "Milímetros",
        "alternativa_e": "Litros",
        "gabarito": "C"
    },
    {
        "codigo": "D6",
        "enunciado": "Qual unidade de medida é mais adequada para medir a distância entre duas cidades?",
        "alternativa_a": "Milímetros",
        "alternativa_b": "Centímetros",
        "alternativa_c": "Metros",
        "alternativa_d": "Quilômetros",
        "alternativa_e": "Gramas",
        "gabarito": "D"
    },
    # D7 - Resolver problemas com unidades de medida
    {
        "codigo": "D7",
        "enunciado": "Uma corda tem 3 metros de comprimento. Quantos centímetros tem essa corda?",
        "alternativa_a": "30 cm",
        "alternativa_b": "300 cm",
        "alternativa_c": "3000 cm",
        "alternativa_d": "3 cm",
        "alternativa_e": "0,3 cm",
        "gabarito": "B"
    },
    {
        "codigo": "D7",
        "enunciado": "Maria comprou 2 kg de arroz. Quantos gramas de arroz ela comprou?",
        "alternativa_a": "20 g",
        "alternativa_b": "200 g",
        "alternativa_c": "2000 g",
        "alternativa_d": "20000 g",
        "alternativa_e": "2 g",
        "gabarito": "C"
    },
    # D8 - Relações entre unidades de tempo
    {
        "codigo": "D8",
        "enunciado": "Quantos minutos tem 2 horas?",
        "alternativa_a": "60 minutos",
        "alternativa_b": "90 minutos",
        "alternativa_c": "100 minutos",
        "alternativa_d": "120 minutos",
        "alternativa_e": "180 minutos",
        "gabarito": "D"
    },
    {
        "codigo": "D8",
        "enunciado": "Um filme tem 1 hora e 30 minutos. Quantos minutos dura o filme?",
        "alternativa_a": "60 minutos",
        "alternativa_b": "90 minutos",
        "alternativa_c": "100 minutos",
        "alternativa_d": "130 minutos",
        "alternativa_e": "150 minutos",
        "gabarito": "B"
    },
    # D9 - Horário de início/término
    {
        "codigo": "D9",
        "enunciado": "A aula começa às 7h30 e termina às 11h30. Quanto tempo dura a aula?",
        "alternativa_a": "3 horas",
        "alternativa_b": "4 horas",
        "alternativa_c": "5 horas",
        "alternativa_d": "6 horas",
        "alternativa_e": "2 horas",
        "gabarito": "B"
    },
    # D10 - Trocas entre cédulas e moedas
    {
        "codigo": "D10",
        "enunciado": "Paulo tem uma nota de R$ 10,00 e quer trocar por moedas de R$ 1,00. Quantas moedas ele receberá?",
        "alternativa_a": "5 moedas",
        "alternativa_b": "10 moedas",
        "alternativa_c": "20 moedas",
        "alternativa_d": "100 moedas",
        "alternativa_e": "1 moeda",
        "gabarito": "B"
    },
    # D11 - Perímetro
    {
        "codigo": "D11",
        "enunciado": "Um quadrado tem lado de 5 cm. Qual é o perímetro desse quadrado?",
        "alternativa_a": "10 cm",
        "alternativa_b": "15 cm",
        "alternativa_c": "20 cm",
        "alternativa_d": "25 cm",
        "alternativa_e": "5 cm",
        "gabarito": "C"
    },
    {
        "codigo": "D11",
        "enunciado": "Um retângulo tem 6 cm de comprimento e 4 cm de largura. Qual é o seu perímetro?",
        "alternativa_a": "10 cm",
        "alternativa_b": "20 cm",
        "alternativa_c": "24 cm",
        "alternativa_d": "14 cm",
        "alternativa_e": "12 cm",
        "gabarito": "B"
    },
    # D12 - Área
    {
        "codigo": "D12",
        "enunciado": "Um quadrado tem lado de 4 cm. Qual é a área desse quadrado?",
        "alternativa_a": "8 cm²",
        "alternativa_b": "12 cm²",
        "alternativa_c": "16 cm²",
        "alternativa_d": "20 cm²",
        "alternativa_e": "4 cm²",
        "gabarito": "C"
    },
    # D13 - Sistema de numeração decimal
    {
        "codigo": "D13",
        "enunciado": "No número 3.456, o algarismo 4 representa:",
        "alternativa_a": "4 unidades",
        "alternativa_b": "4 dezenas",
        "alternativa_c": "4 centenas",
        "alternativa_d": "4 milhares",
        "alternativa_e": "40 unidades",
        "gabarito": "C"
    },
    # D14 - Localização de números na reta numérica
    {
        "codigo": "D14",
        "enunciado": "Na reta numérica: 0 --- 5 --- 10 --- ? --- 20\n\nQual número está faltando?",
        "alternativa_a": "12",
        "alternativa_b": "13",
        "alternativa_c": "14",
        "alternativa_d": "15",
        "alternativa_e": "16",
        "gabarito": "D"
    },
    # D17 - Adição e subtração
    {
        "codigo": "D17",
        "enunciado": "Calcule: 234 + 156 = ?",
        "alternativa_a": "380",
        "alternativa_b": "390",
        "alternativa_c": "400",
        "alternativa_d": "410",
        "alternativa_e": "290",
        "gabarito": "B"
    },
    {
        "codigo": "D17",
        "enunciado": "Calcule: 500 - 237 = ?",
        "alternativa_a": "263",
        "alternativa_b": "273",
        "alternativa_c": "283",
        "alternativa_d": "253",
        "alternativa_e": "363",
        "gabarito": "A"
    },
    # D18 - Multiplicação e divisão
    {
        "codigo": "D18",
        "enunciado": "Calcule: 25 × 4 = ?",
        "alternativa_a": "80",
        "alternativa_b": "90",
        "alternativa_c": "100",
        "alternativa_d": "110",
        "alternativa_e": "29",
        "gabarito": "C"
    },
    {
        "codigo": "D18",
        "enunciado": "Calcule: 144 ÷ 12 = ?",
        "alternativa_a": "10",
        "alternativa_b": "11",
        "alternativa_c": "12",
        "alternativa_d": "13",
        "alternativa_e": "14",
        "gabarito": "C"
    },
    # D19 - Problemas com adição/subtração
    {
        "codigo": "D19",
        "enunciado": "João tinha 45 figurinhas e ganhou mais 28. Quantas figurinhas ele tem agora?",
        "alternativa_a": "63 figurinhas",
        "alternativa_b": "73 figurinhas",
        "alternativa_c": "83 figurinhas",
        "alternativa_d": "17 figurinhas",
        "alternativa_e": "53 figurinhas",
        "gabarito": "B"
    },
    # D20 - Problemas com multiplicação/divisão
    {
        "codigo": "D20",
        "enunciado": "Uma caixa tem 6 lápis. Quantos lápis há em 8 caixas?",
        "alternativa_a": "14 lápis",
        "alternativa_b": "36 lápis",
        "alternativa_c": "42 lápis",
        "alternativa_d": "48 lápis",
        "alternativa_e": "56 lápis",
        "gabarito": "D"
    },
    {
        "codigo": "D20",
        "enunciado": "Maria tem 36 balas para dividir igualmente entre 4 amigos. Quantas balas cada amigo receberá?",
        "alternativa_a": "6 balas",
        "alternativa_b": "7 balas",
        "alternativa_c": "8 balas",
        "alternativa_d": "9 balas",
        "alternativa_e": "10 balas",
        "gabarito": "D"
    },
    # D21 - Representações de números racionais
    {
        "codigo": "D21",
        "enunciado": "Qual das opções representa a fração 1/2 (um meio)?",
        "alternativa_a": "0,25",
        "alternativa_b": "0,50",
        "alternativa_c": "0,75",
        "alternativa_d": "1,00",
        "alternativa_e": "0,20",
        "gabarito": "B"
    },
    # D23 - Escrita decimal de cédulas e moedas
    {
        "codigo": "D23",
        "enunciado": "João tem 2 reais e 50 centavos. Como escrever esse valor em reais?",
        "alternativa_a": "R$ 2,05",
        "alternativa_b": "R$ 2,50",
        "alternativa_c": "R$ 25,00",
        "alternativa_d": "R$ 0,25",
        "alternativa_e": "R$ 250,00",
        "gabarito": "B"
    },
    # D24 - Fração como representação
    {
        "codigo": "D24",
        "enunciado": "Uma pizza foi dividida em 8 partes iguais. Pedro comeu 3 partes. Que fração da pizza Pedro comeu?",
        "alternativa_a": "3/5",
        "alternativa_b": "8/3",
        "alternativa_c": "3/8",
        "alternativa_d": "5/8",
        "alternativa_e": "1/3",
        "gabarito": "C"
    },
    # D26 - Porcentagem
    {
        "codigo": "D26",
        "enunciado": "Em uma sala com 20 alunos, 50% são meninas. Quantas meninas há na sala?",
        "alternativa_a": "5 meninas",
        "alternativa_b": "10 meninas",
        "alternativa_c": "15 meninas",
        "alternativa_d": "20 meninas",
        "alternativa_e": "8 meninas",
        "gabarito": "B"
    },
    {
        "codigo": "D26",
        "enunciado": "Uma loja dá 25% de desconto em um produto de R$ 100,00. Qual é o valor do desconto?",
        "alternativa_a": "R$ 10,00",
        "alternativa_b": "R$ 15,00",
        "alternativa_c": "R$ 20,00",
        "alternativa_d": "R$ 25,00",
        "alternativa_e": "R$ 50,00",
        "gabarito": "D"
    },
    # D27 - Ler informações em tabelas
    {
        "codigo": "D27",
        "enunciado": "Observe a tabela:\n\nFRUTA | QUANTIDADE\nMaçã  | 15\nLaranja | 20\nBanana | 10\n\nQual fruta tem a maior quantidade?",
        "alternativa_a": "Maçã",
        "alternativa_b": "Laranja",
        "alternativa_c": "Banana",
        "alternativa_d": "Todas iguais",
        "alternativa_e": "Não é possível saber",
        "gabarito": "B"
    },
    # D28 - Ler informações em gráficos
    {
        "codigo": "D28",
        "enunciado": "Em um gráfico de barras, a barra do mês de janeiro tem altura 30 e a do mês de fevereiro tem altura 45. Qual mês teve maior valor?",
        "alternativa_a": "Janeiro",
        "alternativa_b": "Fevereiro",
        "alternativa_c": "Os dois iguais",
        "alternativa_d": "Nenhum dos dois",
        "alternativa_e": "Não é possível saber",
        "gabarito": "B"
    },
]

def criar_excel_questoes():
    """Cria arquivo Excel com banco de questões"""
    wb = openpyxl.Workbook()
    
    # Estilos
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    alt_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Criar abas
    ws_port5 = wb.active
    ws_port5.title = "Português 5º Ano"
    
    ws_mat5 = wb.create_sheet("Matemática 5º Ano")
    
    # Cabeçalhos
    headers = ["Código Descritor", "Enunciado", "Alternativa A", "Alternativa B", 
               "Alternativa C", "Alternativa D", "Alternativa E", "Gabarito", "Bloco"]
    
    def preencher_aba(ws, questoes, disciplina, ano):
        # Adicionar cabeçalhos
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        
        # Adicionar questões
        row = 2
        bloco = 1
        questoes_bloco = 0
        
        for q in questoes:
            # Alternar bloco a cada ~11 questões
            if questoes_bloco >= 11:
                bloco = 2
            
            ws.cell(row=row, column=1, value=q["codigo"]).border = border
            ws.cell(row=row, column=2, value=q["enunciado"]).border = border
            ws.cell(row=row, column=3, value=q["alternativa_a"]).border = border
            ws.cell(row=row, column=4, value=q["alternativa_b"]).border = border
            ws.cell(row=row, column=5, value=q["alternativa_c"]).border = border
            ws.cell(row=row, column=6, value=q["alternativa_d"]).border = border
            ws.cell(row=row, column=7, value=q["alternativa_e"]).border = border
            ws.cell(row=row, column=8, value=q["gabarito"]).border = border
            ws.cell(row=row, column=9, value=f"BLOCO_{bloco}").border = border
            
            # Aplicar cor alternada
            if row % 2 == 0:
                for col in range(1, 10):
                    ws.cell(row=row, column=col).fill = alt_fill
            
            # Wrap text para enunciado
            ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
            
            row += 1
            questoes_bloco += 1
        
        # Ajustar larguras das colunas
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 60
        ws.column_dimensions['C'].width = 30
        ws.column_dimensions['D'].width = 30
        ws.column_dimensions['E'].width = 30
        ws.column_dimensions['F'].width = 30
        ws.column_dimensions['G'].width = 30
        ws.column_dimensions['H'].width = 10
        ws.column_dimensions['I'].width = 12
        
        # Altura das linhas
        ws.row_dimensions[1].height = 30
        for r in range(2, row):
            ws.row_dimensions[r].height = 80
    
    # Preencher abas
    preencher_aba(ws_port5, QUESTOES_PORT_5, "PORTUGUES", 5)
    preencher_aba(ws_mat5, QUESTOES_MAT_5, "MATEMATICA", 5)
    
    # Salvar arquivo
    filename = "banco_questoes_saeb.xlsx"
    wb.save(filename)
    print(f"Arquivo '{filename}' criado com sucesso!")
    print(f"- Português 5º Ano: {len(QUESTOES_PORT_5)} questões")
    print(f"- Matemática 5º Ano: {len(QUESTOES_MAT_5)} questões")
    print(f"Total: {len(QUESTOES_PORT_5) + len(QUESTOES_MAT_5)} questões")
    return filename

if __name__ == "__main__":
    criar_excel_questoes()
